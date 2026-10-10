"""Produce an immutable DOCX DFAPTI evidence snapshot from canonical registers."""
import datetime as dt
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / "cases/DFAPTI-NSW-AIS-2026-00001"
DATA = CASE / "data"
DEST = CASE / "snapshots" / "MAYHEM-NSW-AIS-2026-00001-ER-SNAPSHOT-001.docx"
if DEST.exists():
    print("Snapshot already exists; immutable origin retained.")
    raise SystemExit(0)
evidence = json.loads((DATA / "evidence_register.json").read_text())
sources = json.loads((DATA / "source_register.json").read_text())
notes = json.loads((DATA / "investigation_notes.json").read_text())
log = json.loads((DATA / "automation_log.json").read_text())
doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
doc.add_heading("MAYHEM — DFAPTI EVIDENCE REGISTER", 0)
doc.add_heading("SNAPSHOT-001 | Origin record", 1)
doc.add_paragraph("Case: DFAPTI-NSW-AIS-2026-00001")
doc.add_paragraph("Civilian Surveillance Data Transfer and AI Processing — NSW Preventive Governance Audit")
doc.add_paragraph("Status: ACTIVE / NOT FROZEN. No evidence freeze authorised.")
doc.add_paragraph("Origin source: canonical GitHub case registers. Snapshot created on " + dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds") + ".")
doc.add_paragraph("Evidence entries retain their original verification states. Source capture or a URL alone does not verify a claim.")
doc.add_heading("Evidence Register — uncompressed entries", 1)
for e in evidence:
    doc.add_heading(e["evidence_id"], 2)
    for k,v in e.items():
        doc.add_paragraph(k.replace("_"," ").title() + ": " + (json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else str(v) if v is not None else "Not recorded"))
doc.add_heading("Source Register — continuity inventory", 1)
for s in sources:
    doc.add_heading(s["source_id"] + " — " + s["title"], 2)
    for k,v in s.items():
        doc.add_paragraph(k.replace("_"," ").title() + ": " + str(v))
doc.add_heading("Investigation notes — continuity", 1)
for n in notes:
    doc.add_heading(n["note_id"], 2)
    for k,v in n.items():
        doc.add_paragraph(k.replace("_"," ").title() + ": " + str(v))
doc.add_heading("Automation log — continuity", 1)
for r in log:
    doc.add_heading(r["run_id"], 2)
    for k,v in r.items():
        doc.add_paragraph(k.replace("_"," ").title() + ": " + (json.dumps(v,ensure_ascii=False) if isinstance(v,(list,dict)) else str(v)))
doc.add_paragraph("CHAIN CONTROL: Snapshot-002 must explicitly reference this file. No freezing, conclusions or analytical stage advancement is authorised.")
DEST.parent.mkdir(parents=True,exist_ok=True)
doc.save(DEST)
print(DEST)
